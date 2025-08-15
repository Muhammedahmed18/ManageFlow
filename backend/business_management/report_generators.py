"""
Report Generators
================
PDF and Excel report generation functionality.
"""
import io
import csv
from datetime import datetime
from django.http import HttpResponse
from django.db.models import Sum
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import logging
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

logger = logging.getLogger(__name__)

class ReportGenerator:
    """Generate various types of reports in PDF and Excel formats."""
    
    @staticmethod
    def generate_products_report(business_id, format='excel'):
        """Generate products report."""
        from .models import Product
        
        products = Product.objects.filter(business_id=business_id).select_related('category', 'template')
        
        if format == 'excel':
            return ReportGenerator._generate_products_excel(products, business_id)
        elif format == 'csv':
            return ReportGenerator._generate_products_csv(products, business_id)
        else:
            raise ValueError(f"Unsupported format: {format}")
    
    @staticmethod
    def _generate_products_excel(products, business_id):
        """Generate Excel report for products."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products Report"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws['A1'] = f"Products Report - Business ID: {business_id}"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:F1')
        
        # Summary
        ws['A3'] = "Summary"
        ws['A3'].font = Font(bold=True, size=14)
        
        summary_data = [
            ['Total Products', products.count()],
            ['Active Products', products.filter(is_active=True).count()],
            ['Inactive Products', products.filter(is_active=False).count()],
            ['Report Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')]
        ]
        
        for i, (label, value) in enumerate(summary_data, start=4):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = Font(bold=True)
        
        # Products table
        if products.exists():
            headers = ['ID', 'Name', 'Custom ID', 'Category', 'Status', 'Created Date']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=8, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
            
            for row, product in enumerate(products, start=9):
                ws.cell(row=row, column=1, value=product.id).border = border
                ws.cell(row=row, column=2, value=product.name).border = border
                ws.cell(row=row, column=3, value=product.custom_id or 'N/A').border = border
                ws.cell(row=row, column=4, value=product.category.name if product.category else 'N/A').border = border
                ws.cell(row=row, column=5, value='Active' if product.is_active else 'Inactive').border = border
                ws.cell(row=row, column=6, value=product.created_at.strftime('%Y-%m-%d')).border = border
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="products_report_{business_id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_products_csv(products, business_id):
        """Generate CSV report for products."""
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="products_report_{business_id}_{datetime.now().strftime("%Y%m%d")}.csv"'
        
        writer = csv.writer(response)
        writer.writerow(['ID', 'Name', 'Custom ID', 'Category', 'Status', 'Created Date'])
        
        for product in products:
            writer.writerow([
                product.id,
                product.name,
                product.custom_id or 'N/A',
                product.category.name if product.category else 'N/A',
                'Active' if product.is_active else 'Inactive',
                product.created_at.strftime('%Y-%m-%d')
            ])
        
        return response
    
    @staticmethod
    def generate_orders_report(business_id, format='excel', start_date=None, end_date=None):
        """Generate orders report."""
        from .models import Order
        
        orders = Order.objects.filter(business_id=business_id)
        
        if start_date:
            orders = orders.filter(created_at__gte=start_date)
        if end_date:
            orders = orders.filter(created_at__lte=end_date)
        
        orders = orders.select_related('customer').prefetch_related('items__product')
        
        if format == 'excel':
            return ReportGenerator._generate_orders_excel(orders, business_id)
        elif format == 'csv':
            return ReportGenerator._generate_orders_csv(orders, business_id)
        else:
            raise ValueError(f"Unsupported format: {format}")
    
    @staticmethod
    def _generate_orders_excel(orders, business_id):
        """Generate Excel report for orders."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders Report"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws['A1'] = f"Orders Report - Business ID: {business_id}"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:F1')
        
        # Summary
        total_orders = orders.count()
        total_revenue = orders.aggregate(total=Sum('total_amount'))['total'] or 0
        completed_orders = orders.filter(status='completed').count()
        
        ws['A3'] = "Summary"
        ws['A3'].font = Font(bold=True, size=14)
        
        summary_data = [
            ['Total Orders', total_orders],
            ['Total Revenue', f"${total_revenue:.2f}"],
            ['Completed Orders', completed_orders],
            ['Completion Rate', f"{(completed_orders/total_orders*100):.1f}%" if total_orders > 0 else "0%"],
            ['Report Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')]
        ]
        
        for i, (label, value) in enumerate(summary_data, start=4):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = Font(bold=True)
        
        # Orders table
        if orders.exists():
            headers = ['Order #', 'Customer', 'Status', 'Total Amount', 'Items', 'Created Date']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=9, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
            
            for row, order in enumerate(orders, start=10):
                ws.cell(row=row, column=1, value=order.order_number).border = border
                ws.cell(row=row, column=2, value=order.customer.name if order.customer else 'N/A').border = border
                ws.cell(row=row, column=3, value=order.status.title()).border = border
                ws.cell(row=row, column=4, value=f"${order.total_amount:.2f}").border = border
                ws.cell(row=row, column=5, value=order.items.count()).border = border
                ws.cell(row=row, column=6, value=order.created_at.strftime('%Y-%m-%d')).border = border
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="orders_report_{business_id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_orders_csv(orders, business_id):
        """Generate CSV report for orders."""
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="orders_report_{business_id}_{datetime.now().strftime("%Y%m%d")}.csv"'
        
        writer = csv.writer(response)
        writer.writerow(['Order #', 'Customer', 'Status', 'Total Amount', 'Items', 'Created Date'])
        
        for order in orders:
            writer.writerow([
                order.order_number,
                order.customer.name if order.customer else 'N/A',
                order.status.title(),
                f"${order.total_amount:.2f}",
                order.items.count(),
                order.created_at.strftime('%Y-%m-%d')
            ])
        
        return response

    @staticmethod
    def generate_invoice_pdf(invoice):
        """Generate PDF invoice."""
        try:
            print(f"ReportGenerator: Generating PDF for invoice {invoice.id}")
            print(f"ReportGenerator: Invoice number: {invoice.invoice_number}")
            print(f"ReportGenerator: Invoice type: {invoice.invoice_type}")
            print(f"ReportGenerator: Line items: {invoice.line_items}")
            print(f"ReportGenerator: Has items: {hasattr(invoice, 'items')}")
            if hasattr(invoice, 'items'):
                print(f"ReportGenerator: Items count: {invoice.items.count()}")
            
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            elements = []
            
            # Get styles
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                spaceAfter=30,
                alignment=TA_CENTER
            )
            heading_style = ParagraphStyle(
                'CustomHeading',
                parent=styles['Heading2'],
                fontSize=14,
                spaceAfter=12,
                spaceBefore=12
            )
            normal_style = styles['Normal']
            
            # Title
            elements.append(Paragraph(f"INVOICE #{invoice.invoice_number or 'N/A'}", title_style))
            elements.append(Spacer(1, 20))
            
            # Invoice details table
            invoice_data = [
                ['Invoice Number:', invoice.invoice_number or 'N/A'],
                ['Invoice Date:', invoice.invoice_date.strftime('%B %d, %Y') if invoice.invoice_date else 'N/A'],
                ['Due Date:', invoice.due_date.strftime('%B %d, %Y') if invoice.due_date else 'N/A'],
                ['Status:', invoice.status.title() if invoice.status else 'N/A'],
            ]
            
            if invoice.recipient_name:
                invoice_data.append(['Bill To:', invoice.recipient_name])
            if invoice.bill_to:
                invoice_data.append(['Address:', invoice.bill_to])
            if invoice.contact_info:
                invoice_data.append(['Contact:', invoice.contact_info])
            
            invoice_table = Table(invoice_data, colWidths=[2*inch, 4*inch])
            invoice_table.setStyle(TableStyle([
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(invoice_table)
            elements.append(Spacer(1, 20))
            
            # Line items table
            if hasattr(invoice, 'items') and invoice.items.exists():
                # Use InvoiceItem model
                items_data = [['Item', 'Description', 'Qty', 'Unit Price', 'Total']]
                for item in invoice.items.all():
                    items_data.append([
                        item.product_name or 'Product',
                        item.description or '',
                        str(item.quantity or 1),
                        f"${float(item.unit_price or 0):.2f}",
                        f"${float(item.total_price or 0):.2f}"
                    ])
            elif invoice.line_items:
                # Use JSON line_items field
                items_data = [['Item', 'Description', 'Qty', 'Unit Price', 'Total']]
                for item in invoice.line_items:
                    items_data.append([
                        item.get('product_name', 'Product'),
                        item.get('description', ''),
                        str(item.get('quantity', 1)),
                        f"${float(item.get('unit_price', 0)):.2f}",
                        f"${float(item.get('amount', 0)):.2f}"
                    ])
            else:
                items_data = [['Item', 'Description', 'Qty', 'Unit Price', 'Total']]
                # Add a placeholder row if no items
                items_data.append(['No items', '', '0', '$0.00', '$0.00'])
            
            items_table = Table(items_data, colWidths=[1.5*inch, 2*inch, 0.8*inch, 1.2*inch, 1.2*inch])
            items_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('ALIGN', (2, 1), (2, -1), 'CENTER'),  # Quantity
                ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),  # Prices
            ]))
            elements.append(items_table)
            elements.append(Spacer(1, 20))
            
            # Totals
            totals_data = [
                ['Subtotal:', f"${float(invoice.subtotal or 0):.2f}"],
                ['Tax ({:.1f}%):'.format(float(invoice.tax_percentage or 0)), f"${float(invoice.tax_amount or 0):.2f}"],
                ['Total Amount:', f"${float(invoice.total_amount or 0):.2f}"],
            ]
            
            if invoice.advanced_paid and float(invoice.advanced_paid) > 0:
                totals_data.append(['Advanced Paid:', f"${float(invoice.advanced_paid):.2f}"])
                totals_data.append(['Balance Due:', f"${float(invoice.balance_due or 0):.2f}"])
            
            totals_table = Table(totals_data, colWidths=[4*inch, 2*inch])
            totals_table.setStyle(TableStyle([
                ('ALIGN', (0, 0), (0, -1), 'RIGHT'),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(totals_table)
            elements.append(Spacer(1, 30))
            
            # Notes
            if invoice.notes:
                elements.append(Paragraph("Notes:", heading_style))
                elements.append(Paragraph(invoice.notes, normal_style))
                elements.append(Spacer(1, 20))
            
            # Payment instructions
            if invoice.payment_instructions:
                elements.append(Paragraph("Payment Instructions:", heading_style))
                elements.append(Paragraph(invoice.payment_instructions, normal_style))
                elements.append(Spacer(1, 20))
            
            # Thank you message
            if invoice.thank_you_message:
                elements.append(Paragraph(invoice.thank_you_message, normal_style))
            
            # Build PDF
            doc.build(elements)
            buffer.seek(0)
            
            return buffer
            
        except Exception as e:
            logger.error(f"Error generating PDF for invoice {invoice.id}: {str(e)}")
            print(f"Error generating PDF for invoice {invoice.id}: {str(e)}")
            raise e